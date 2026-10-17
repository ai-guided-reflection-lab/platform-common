"""Acceptance checks for the professor-configured, persistent student conversation."""
import copy
import importlib
import json
from pathlib import Path
from uuid import uuid4

import httpx
import pytest
from fastapi.testclient import TestClient


def headers(user="student-alex"):
    return {"X-Demo-User": user}


def lesson(name="requirements"):
    return json.loads((Path(__file__).parents[1] / "lessons" / f"{name}.json").read_text(encoding="utf-8"))


@pytest.fixture
def client(tmp_path, monkeypatch):
    for key, value in {"ADAPTIVE_DATA_DIR": str(tmp_path), "ADAPTIVE_DATABASE_PATH": str(tmp_path / "test.sqlite3"), "ADAPTIVE_UPLOAD_DIR": str(tmp_path / "uploads"), "ADAPTIVE_AI_MODE": "demo", "OPENAI_API_KEY": "", "GROQ_API_KEY": ""}.items():
        monkeypatch.setenv(key, value)
    from app import config, db, main, openai_client, rag, lesson_flow
    for module in (config, openai_client, db, rag, lesson_flow, main):
        importlib.reload(module)
    with TestClient(main.app) as test_client:
        yield test_client


def test_document_preview_reads_material_without_creating_assignment(client):
    before = client.get("/api/instructor/assignments", headers=headers("instructor-demo")).json()
    response = client.post("/api/instructor/documents/preview", headers=headers("instructor-demo"), files={"file": ("course.md", b"# Course material\nReason through an example.", "text/markdown")})
    assert response.status_code == 200, response.text
    assert response.json() == {"filename": "course.md", "text": "# Course material\nReason through an example.", "truncated": False}
    assert client.get("/api/instructor/assignments", headers=headers("instructor-demo")).json() == before
    assert client.post("/api/instructor/documents/preview", headers=headers(), files={"file": ("course.txt", b"Material")}).status_code == 403


def test_document_preview_limits_and_unreadable_files(client):
    def upload(name, content):
        return client.post("/api/instructor/documents/preview", headers=headers("instructor-demo"), files={"file": (name, content)})
    assert upload("course.docx", b"content").status_code == 422
    assert upload("course.txt", b"   ").status_code == 422
    assert upload("course.pdf", b"invalid PDF").status_code == 422
    assert upload("course.txt", b"x" * (10 * 1024 * 1024 + 1)).status_code == 413
    response = upload("course.txt", b"x" * 30001)
    assert response.status_code == 200
    assert response.json()["truncated"] is True
    assert len(response.json()["text"]) == 30000


def test_word_document_preview_extracts_paragraphs_and_tables(client):
    from io import BytesIO
    from docx import Document
    document = Document()
    document.add_paragraph("Students reason through an example.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Concept"
    table.cell(0, 1).text = "Evidence"
    content = BytesIO()
    document.save(content)
    response = client.post("/api/instructor/documents/preview", headers=headers("instructor-demo"), files={"file": ("course.docx", content.getvalue())})
    assert response.status_code == 200, response.text
    assert "Students reason through an example." in response.json()["text"]
    assert "Concept | Evidence" in response.json()["text"]


def test_demo_document_indexing_does_not_contact_ai_even_with_a_key(client, monkeypatch):
    from app import openai_client
    monkeypatch.setattr(openai_client, "OPENAI_API_KEY", "unused-test-key")
    def unexpected(*args, **kwargs):
        raise AssertionError("Demo indexing must not contact the AI provider")
    monkeypatch.setattr(httpx, "Client", unexpected)
    aid = create(client, publish=False)
    response = client.post(f"/api/instructor/assignments/{aid}/documents", headers=headers("instructor-demo"), files={"file": ("course.txt", b"Students learn from examples and explain their reasoning.")})
    assert response.status_code == 201, response.text
    assert response.json()["chunk_count"] == 1


def test_document_review_uses_material_and_prepares_unpublished_lesson(client, monkeypatch):
    from app import main
    expected = lesson()
    seen = {}
    def infer(material, filename):
        seen.update(material=material, filename=filename)
        return {"topic": expected["topic"], "objective_descriptions": [o["description"] for o in expected["objectives"]]}
    monkeypatch.setattr(main, "infer_document_context", infer)
    monkeypatch.setattr(main, "discover_resources", lambda topic: expected["study_resources"])
    monkeypatch.setattr(main, "verify_resources", lambda candidates, topic: candidates)
    monkeypatch.setattr(main, "read_resource_excerpts", lambda resources: [])
    monkeypatch.setattr(main, "generate_learning_plan", lambda *args, **kwargs: {**expected, "content_approved": False})
    aid = create(client, publish=False)
    response = client.post(f"/api/instructor/assignments/{aid}/prepare-review", headers=headers("instructor-demo"))
    assert response.status_code == 200, response.text
    assert seen["material"] == expected["approved_material"]
    assert response.json()["status"] == "draft"
    assert response.json()["learning_plan"]["content_approved"] is False
    assert client.post("/api/instructor/document-draft", headers=headers(), json={"material": expected["approved_material"]}).status_code == 403


@pytest.mark.parametrize("provider_body", [
    {"error": {"message": "Request too large for organization org_private: Limit 8000, Requested 11001", "code": "rate_limit_exceeded"}},
    {"error": None},
    {"error": {"message": None}},
])
def test_document_review_explains_provider_size_limit(client, monkeypatch, provider_body):
    from app import main
    def fail(*args):
        response = httpx.Response(413, json=provider_body, request=httpx.Request("POST", "https://api.groq.com/openai/v1/chat/completions"))
        response.raise_for_status()
    monkeypatch.setattr(main, "infer_document_context", fail)
    response = client.post("/api/instructor/document-draft", headers=headers("instructor-demo"), json={"filename": "course.txt", "material": "A course document with enough material for review."})
    assert response.status_code == 502
    detail = response.json()["detail"]
    assert "too large" in detail
    assert "Retrying the same document" in detail
    assert "org_private" not in detail
    if provider_body.get("error") and provider_body["error"].get("message"):
        assert "Token allowance: 8000; requested: 11001" in detail


def test_failed_document_review_preserves_draft(client, monkeypatch):
    from app import main
    aid = create(client, publish=False)
    before = client.get(f"/api/instructor/assignments/{aid}", headers=headers("instructor-demo")).json()
    monkeypatch.setattr(main, "infer_document_context", lambda *args: {"topic": "Course topic", "objective_descriptions": ["Explain an example."]})
    monkeypatch.setattr(main, "discover_resources", lambda topic: [])
    monkeypatch.setattr(main, "verify_resources", lambda *args: [])
    response = client.post(f"/api/instructor/assignments/{aid}/prepare-review", headers=headers("instructor-demo"))
    assert response.status_code == 502
    assert "Credible readings" in response.json()["detail"]
    assert client.get(f"/api/instructor/assignments/{aid}", headers=headers("instructor-demo")).json() == before


def test_generated_practice_repairs_map_to_original_quiz_concepts(client, monkeypatch):
    from app import openai_client
    plan = lesson()
    banks = {}
    for asset in plan["learning_assets"]:
        activities = {}
        for question in plan["diagnostic_quiz"]:
            if question["objective_id"] == asset["objective_id"]:
                entry = {"quiz_question_id": question["id"]}
                for level in ("foundational", "standard", "accelerated"):
                    entry[level] = {**asset[level][0], "concept": "An incorrect label from the model", "scenario": "Consider a fresh course example with a different concrete situation."}
                activities[question["id"] + "-case-1"] = entry
        banks[asset["objective_id"]] = activities
    repair = {"question_repairs": {q["id"]: {"distractor_misconceptions": q["distractor_misconceptions"], "reasoning_rubric": q["reasoning_rubric"]} for q in plan["diagnostic_quiz"]}, "practice_banks": banks}
    monkeypatch.setattr(openai_client, "_chat_settings", lambda: ("https://api.openai.com/v1", "test-model", {}))
    def repair_response(*args, **kwargs):
        prompt = kwargs['json']['messages'][0]['content']
        assert "refer to concrete details from that question's scenario" in prompt
        assert 'ONE actionable next step or comparison in plain language' in prompt
        assert 'leave the final decision and explanation to them' in prompt
        assert 'pasted unchanged into an unrelated question' in prompt
        assert 'generate reasoning' in prompt
        assert 'not merely entity names' in prompt
        assert 'Do not state the conclusion' in prompt
        assert 'not exact wording or a single keyword' in prompt
        return httpx.Response(200, request=httpx.Request("POST", "https://api.openai.com/v1/chat/completions"), json={"choices": [{"message": {"content": json.dumps(repair)}}]})
    monkeypatch.setattr(httpx.Client, "post", repair_response)
    result = openai_client.repair_learning_plan(plan, ["Map practice to diagnostic concepts."])
    assert result["content_approved"] is False
    for asset in result["learning_assets"]:
        expected = {q["concept"] for q in plan["diagnostic_quiz"] if q["objective_id"] == asset["objective_id"]}
        for level in ("foundational", "standard", "accelerated"):
            assert {activity["concept"] for activity in asset[level]} == expected


def create(client, plan=None, publish=True, students=None):
    plan = plan or lesson()
    response = client.post("/api/instructor/assignments", headers=headers("instructor-demo"), json={"title": plan["title"], "instructions": "Study and apply the approved material.", "learning_plan": plan, "student_ids": students or ["student-alex", "student-jordan"]})
    assert response.status_code == 201, response.text
    aid = response.json()["id"]
    if publish:
        response = client.post(f"/api/instructor/assignments/{aid}/publish", headers=headers("instructor-demo"))
        assert response.status_code == 200, response.text
    return aid


def turn(client, aid, content="", user="student-alex", **fields):
    response = client.post(f"/api/student/assignments/{aid}/learning-turn", headers=headers(user), json={"turn_id": str(uuid4()), "content": content, **fields})
    assert response.status_code == 200, response.text
    return response.json()


def start_quiz(client, aid, user="student-alex"):
    response = client.post(f"/api/student/assignments/{aid}/start", headers=headers(user))
    assert response.status_code == 200, response.text
    assert response.json()["phase"] == "welcome"
    assert "What feels most unfamiliar" in response.json()["messages"][-1]["content"]
    studied = turn(client, aid, "I am completely new to this", user)
    assert studied["phase"] == "study_resources"
    assert studied["learning_state"]["current_paths"] == {}
    assert turn(client, aid, "I'm not ready yet", user)["phase"] == "study_resources"
    return turn(client, aid, "I've finished studying, let's continue", user)


def reasoning(plan, index):
    if plan["topic"] == "Requirements Engineering":
        return "The limits conflict, so we should reconcile them with stakeholders." if index == 3 else "Check feasibility and cost trade-offs before promising a reliability target."
    return "Multiply the numerator and denominator by the same number to preserve the proportion."


def quiz(client, aid, plan, answers=None, confidence="high", explanations=True, user="student-alex"):
    answers = answers if answers is not None else [q["correct_index"] for q in plan["diagnostic_quiz"]]
    for index, (q, option) in enumerate(zip(plan["diagnostic_quiz"], answers)):
        result = turn(client, aid, "I don't know" if option is None else "", user, action="answer", question_id=q["id"], option_index=option, confidence=confidence, explanation=reasoning(plan, index) if index >= 3 and explanations else None)
        if result["learning_state"].get("metadata_requested"):
            result = turn(client, aid, action="continue", user=user)
        if index < 4:
            assert "correct" not in result["learning_state"]["diagnostic"][-1]
            assert "explanation_of_answer" not in result["learning_state"]["diagnostic"][-1]
    return result


@pytest.mark.parametrize("name", ["requirements", "fractions"])
def test_two_topics_common_quiz_fresh_evidence_reflection_and_open_session(client, name):
    plan = lesson(name)
    aid = create(client, plan)
    first = start_quiz(client, aid)
    second = start_quiz(client, aid, "student-jordan")
    assert first["learning_state"]["current_question"] == second["learning_state"]["current_question"]
    detail = client.get(f"/api/student/assignments/{aid}", headers=headers()).json()
    assert detail["learning_plan"]["diagnostic_quiz"] == []
    assert "lesson_snapshot" not in json.dumps(first)
    assert "correct_index" not in json.dumps(first)
    completed = quiz(client, aid, plan)
    assert completed["quiz_score"] == 5
    assert completed["learning_state"]["selected_followup"] == "application_challenge"
    assert completed["phase"] == "learning_choice"
    assert len(completed["learning_state"]["choices"]) == 5
    completed = turn(client, aid, "An application challenge", action="choice")
    assert completed["learning_state"]["choices"] == []
    assert completed["learning_state"]["current_activity"]["prompt"] not in [q["question"] for q in plan["diagnostic_quiz"]]
    original = copy.deepcopy(completed["learning_state"]["diagnostic"])
    response = "Sending a ticket is functional behavior, whereas delivering it within ten seconds is a non-functional quality constraint." if name == "requirements" else "The denominator counts all equal parts of the whole, including selected and unselected parts."
    demonstrated = turn(client, aid, response)
    assert demonstrated["phase"] == "adaptive_learning"
    assert demonstrated["learning_state"]["choices"] == []
    deeper = "The confirmation action is functional behavior, while the deadline is a non-functional quality constraint." if name == "requirements" else "The denominator gives the total number of equal parts making up the whole ribbon."
    demonstrated = turn(client, aid, deeper)
    assert demonstrated["phase"] == "reflection"
    assert "What can you explain or do now" in demonstrated["messages"][-1]["content"]
    reflected = turn(client, aid, "I can apply the idea and justify a new example now.")
    assert reflected["phase"] == "learning_choice"
    assert reflected["status"] != "completed"
    assert "Would you like another example" in reflected["messages"][-1]["content"]
    assert reflected["learning_state"]["diagnostic"] == original
    assert turn(client, aid, action="pause")["phase"] == "paused"
    assert turn(client, aid, action="resume")["phase"] == "learning_choice"
    assert turn(client, aid, "A harder challenge", action="choice")["phase"] == "adaptive_learning"
    evidence = client.get(f"/api/instructor/assignments/{aid}/students/student-alex", headers=headers("instructor-demo"))
    assert evidence.status_code == 200
    assert evidence.json()["attempt"]["learning_state"]["reflections"]


@pytest.mark.parametrize("case,expected", [("foundational", "simpler_support"), ("application", "targeted_practice"), ("confident_wrong", "counterexample_revision"), ("unclear_reasoning", "reasoning_probe"), ("low_confidence", "confidence_building"), ("unknown", "clarification_choice")])
def test_diagnostic_adapts_to_evidence_without_weighting_confidence(client, case, expected):
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    answers = [q["correct_index"] for q in plan["diagnostic_quiz"]]
    confidence, explanations = "medium", True
    if case in {"foundational", "application", "confident_wrong"}:
        index = 2 if case == "application" else 0
        answers[index] = (answers[index] + 1) % 4
        confidence = "high" if case == "confident_wrong" else "medium"
    elif case == "unclear_reasoning":
        explanations = False
    elif case == "low_confidence":
        confidence = "low"
    else:
        answers = [None] * 5
    result = quiz(client, aid, plan, answers, confidence, explanations)
    assert result["learning_state"]["selected_followup"] == expected
    assert result["quiz_score"] == sum(a == q["correct_index"] for a, q in zip(answers, plan["diagnostic_quiz"]))
    assert result["learning_state"]["demonstrated"] == []
    if case == "confident_wrong":
        assert result["phase"] == "learning_choice"
        selected = turn(client, aid, "A simpler explanation", action="choice")
        assert "earlier choice" in selected["messages"][-1]["content"].lower()


def test_missing_metadata_once_hints_pause_ambiguity_and_idempotent_retry(client):
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    q = plan["diagnostic_quiz"][0]
    assert turn(client, aid, "A or B")["learning_state"]["quiz_index"] == 0
    hinted = turn(client, aid, action="hint")
    assert q["options"][q["correct_index"]] not in hinted["messages"][-1]["content"]
    paused = turn(client, aid, action="pause")
    assert paused["learning_state"]["current_question"]["id"] == q["id"]
    turn(client, aid, action="resume")
    pending = turn(client, aid, "I think option A")
    assert pending["learning_state"]["metadata_requested"]
    assert client.get(f"/api/student/assignments/{aid}/attempt", headers=headers()).json()["learning_state"]["metadata_requested"]
    continued = turn(client, aid, action="continue", option_index=3)
    recorded = continued["learning_state"]["diagnostic"][0]
    assert recorded["option_index"] == 0
    assert recorded["missing_confidence"] and recorded["assistance_used"]
    q = plan["diagnostic_quiz"][1]
    body = {"turn_id": "same-request", "action": "answer", "content": q["options"][q["correct_index"]], "confidence": "high", "question_id": q["id"]}
    url = f"/api/student/assignments/{aid}/learning-turn"
    first = client.post(url, headers=headers(), json=body)
    assert first.status_code == 200, first.text
    assert client.post(url, headers=headers(), json=body).json() == first.json()
    assert client.post(url, headers=headers(), json={**body, "confidence": "low"}).status_code == 409
    assert client.post(url, headers=headers(), json={**body, "turn_id": "stale"}).status_code == 409


def test_lesson_turn_save_rolls_back_and_same_turn_can_be_retried(client, monkeypatch):
    from app import db
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    quiz(client, aid, plan)
    turn(client, aid, "Another practice question", action="choice")
    with db.connection() as conn:
        before = dict(conn.execute("SELECT * FROM attempts WHERE assignment_id=? AND student_id=?", (aid, "student-alex")).fetchone())
        message_count = conn.execute("SELECT count(*) AS n FROM messages WHERE attempt_id=?", (before["id"],)).fetchone()["n"]
    original_execute = db.DatabaseConnection.execute
    def fail_evidence_insert(self, statement, parameters=()):
        if statement.startswith("INSERT INTO evidence"):
            raise RuntimeError("Simulated persistence failure")
        return original_execute(self, statement, parameters)
    body = {"turn_id": "retry-after-rollback", "content": "A ticket is functional behavior and the delivery deadline is a non-functional quality constraint."}
    url = f"/api/student/assignments/{aid}/learning-turn"
    with monkeypatch.context() as patch:
        patch.setattr(db.DatabaseConnection, "execute", fail_evidence_insert)
        with pytest.raises(RuntimeError, match="Simulated persistence failure"):
            client.post(url, headers=headers(), json=body)
    with db.connection() as conn:
        assert dict(conn.execute("SELECT * FROM attempts WHERE id=?", (before["id"],)).fetchone()) == before
        assert conn.execute("SELECT count(*) AS n FROM messages WHERE attempt_id=?", (before["id"],)).fetchone()["n"] == message_count
        assert conn.execute("SELECT count(*) AS n FROM evidence WHERE attempt_id=?", (before["id"],)).fetchone()["n"] == 0
    retried = client.post(url, headers=headers(), json=body)
    assert retried.status_code == 200, retried.text
    assert client.post(url, headers=headers(), json=body).json() == retried.json()
    with db.connection() as conn:
        assert conn.execute("SELECT count(*) AS n FROM messages WHERE attempt_id=?", (before["id"],)).fetchone()["n"] == message_count + 2
        assert conn.execute("SELECT count(*) AS n FROM evidence WHERE attempt_id=?", (before["id"],)).fetchone()["n"] == 1


@pytest.mark.parametrize("postgres", [False, True])
def test_turn_lookup_uses_the_callers_connection_and_preserves_locking(postgres):
    from app.db import load_attempt_for_turn
    from types import SimpleNamespace
    calls = []
    def execute(statement, parameters):
        calls.append((statement, parameters))
        return SimpleNamespace(fetchone=lambda: {"id": "attempt"})
    conn = SimpleNamespace(postgres=postgres, execute=execute)
    assert load_attempt_for_turn(conn, "assignment", "student") == {"id": "attempt"}
    assert calls == [("SELECT * FROM attempts WHERE assignment_id=? AND student_id=?" + (" FOR UPDATE" if postgres else ""), ("assignment", "student"))]


def test_hints_require_fresh_independent_check_and_preserve_original_results(client):
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    answers = [q["correct_index"] for q in plan["diagnostic_quiz"]]
    answers[0] = (answers[0] + 1) % 4
    before = quiz(client, aid, plan, answers, "medium")
    original = before["learning_state"]["diagnostic"]
    initial_paths = before["learning_state"]["diagnostic_paths"]
    turn(client, aid, "Another practice question", action="choice")
    turn(client, aid, action="hint")
    assisted = turn(client, aid, "A ticket is functional behavior and the delivery deadline is a non-functional quality constraint.")
    assert assisted["phase"] == "adaptive_learning"
    assert assisted["learning_state"]["demonstrated"] == []
    assert not assisted["learning_state"]["current_activity"]["assisted"]
    fresh = turn(client, aid, "A booking action is functional behavior; response time is a non-functional quality constraint.")
    assert fresh["phase"] == "adaptive_learning"
    fresh = turn(client, aid, "Renewing the loan is a functional action, whereas its time limit is a non-functional quality constraint.")
    assert fresh["phase"] == "reflection"
    assert fresh["learning_state"]["current_paths"] != initial_paths
    assert fresh["learning_state"]["diagnostic_paths"] == initial_paths
    assert fresh["learning_state"]["diagnostic"] == original


def test_publishing_flags_missing_content_and_does_not_fabricate(client):
    plan = lesson()
    plan["diagnostic_quiz"], plan["approved_material"] = [], ""
    aid = create(client, plan, publish=False)
    detail = client.get(f"/api/instructor/assignments/{aid}", headers=headers("instructor-demo")).json()
    assert detail["configuration_issues"]
    assert detail["learning_plan"]["diagnostic_quiz"] == []
    assert client.post(f"/api/instructor/assignments/{aid}/publish", headers=headers("instructor-demo")).status_code == 422
    assert client.post(f"/api/student/assignments/{aid}/start", headers=headers()).status_code == 404
    assert client.post("/api/instructor/generate-plan", headers=headers("instructor-demo"), json={"topic": "Any subject"}).status_code == 422


def test_uploaded_course_notes_retrieved_and_published_configuration_frozen(client):
    plan = lesson()
    aid = create(client, plan, publish=False)
    url = f"/api/instructor/assignments/{aid}"
    upload = client.post(url + "/documents", headers=headers("instructor-demo"), files={"file": ("requirements-notes.txt", b"Functional behavior is an action. A non-functional quality constraint specifies response time.", "text/plain")})
    assert upload.status_code == 201, upload.text
    assert upload.json()["chunk_count"] == 1
    assert client.post(url + "/publish", headers=headers("instructor-demo")).status_code == 200
    client.post(f"/api/student/assignments/{aid}/start", headers=headers())
    answer = turn(client, aid, "Can you explain functional behavior using course notes?")
    assert answer["messages"][-1]["sources"][0]["filename"] == "requirements-notes.txt"
    assert client.put(url, headers=headers("instructor-demo"), json={"title": plan["title"], "learning_plan": plan, "student_ids": ["student-alex"]}).status_code == 409
    assert client.post(url + "/refresh-resources", headers=headers("instructor-demo")).status_code == 409
    assert client.post(url + "/documents", headers=headers("instructor-demo"), files={"file": ("new.txt", b"new material", "text/plain")}).status_code == 409


def test_authorization_and_legacy_endpoints_cannot_skip_stages(client):
    aid = create(client, students=["student-alex"])
    url = f"/api/student/assignments/{aid}"
    assert client.get(url).status_code == 401
    assert client.get(url, headers=headers("student-jordan")).status_code == 404
    assert client.get("/api/instructor/assignments", headers=headers()).status_code == 403
    client.post(url + "/start", headers=headers())
    for suffix, body in [("/quiz", {"answers": [0] * 5}), ("/study-complete", {}), ("/messages", {"content": "Skip"}), ("/submit-task", {"content": "Done"})]:
        assert client.post(url + suffix, headers=headers(), json=body).status_code == 409


def test_provider_failure_uses_configured_rubric_and_disagreement_is_evidence(client, monkeypatch):
    from app import openai_client
    plan = lesson()
    objective = plan["objectives"][0]
    monkeypatch.setattr(openai_client, "runtime_ai_mode", lambda: "openai")
    monkeypatch.setattr(openai_client, "ALLOW_DEMO_FALLBACK", True)
    def unauthorized(*args, **kwargs):
        return httpx.Response(401, request=httpx.Request("POST", "https://api.openai.com/v1/chat/completions"), json={"error": "invalid key"})
    monkeypatch.setattr(httpx.Client, "post", unauthorized)
    assessed = openai_client.assess_configured_response("I disagree that all requirements are functional: a functional action is behavior, whereas a non-functional deadline is a quality constraint.", objective["rubric"], objective, plan["approved_material"], "Classify an action and deadline")
    assert assessed["category"] == "sound_reasoning"
    assert assessed["source"] == "demo_rubric"
    assert assessed["fallback_reason"] == "grader_unavailable"
    assert openai_client.assess_configured_response("Lots of eloquent words without evidence", objective["rubric"], objective, plan["approved_material"], "Classify")["category"] == "insufficient_evidence"


def test_grader_outage_is_not_persisted_as_scored_student_failure(client, monkeypatch):
    from app import lesson_flow, openai_client, db
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    quiz(client, aid, plan)
    turn(client, aid, "Another practice question", action="choice")
    monkeypatch.setattr(lesson_flow, "assess_configured_response", lambda *args: openai_client._assessment_issue("grader_unavailable"))
    result = turn(client, aid, "A ticket is functional behavior and the deadline is a non-functional quality constraint.")
    assert result["phase"] == "learning_choice"
    assert result["learning_state"]["practice_history"][-1]["category"] == "grader_unavailable"
    assert result["learning_state"]["practice_history"][-1]["score"] is None
    assert result["learning_state"]["demonstrated"] == []
    with db.connection() as conn:
        assert conn.execute("SELECT count(*) AS n FROM evidence").fetchone()["n"] == 0


def test_diagnostic_outage_preserves_correct_mcq_results(client, monkeypatch):
    from app import lesson_flow, openai_client
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    monkeypatch.setattr(lesson_flow, "assess_configured_response", lambda *args: openai_client._assessment_issue("grader_error"))
    result = quiz(client, aid, plan)
    assert result["quiz_score"] == 5
    assert all(d["correct"] for d in result["learning_state"]["diagnostic"])
    assert result["learning_state"]["diagnostic"][3]["need"] == "independent_check"
    assert result["learning_state"]["diagnostic"][3]["reasoning"]["status"] == "grader_error"
    assert "could not be graded" in result["messages"][-1]["content"]


def test_corrected_memory_lesson_new_sessions_leave_old_snapshots_intact(client):
    from app import db
    corrected = lesson("operating-system-processes")
    old = copy.deepcopy(corrected)
    old["content_approved"] = True
    activity = next(a for a in old["learning_assets"][0]["foundational"] if a["concept"] == "Process Memory Layout")
    activity["rubric"] = [{"description": "Identifies network buffer as not part of process memory layout.",
                          "evidence_groups": [["code, data, heap are process memory", "network buffer is hardware/system resource"]]}]
    old_aid = create(client, old)
    assert client.post(f"/api/student/assignments/{old_aid}/start", headers=headers()).status_code == 200
    with db.connection() as conn:
        before = json.loads(conn.execute("SELECT learning_state FROM attempts WHERE assignment_id=?", (old_aid,)).fetchone()["learning_state"])["lesson_snapshot"]
    corrected["content_approved"] = True  # Instructor approval for this isolated test assignment.
    new_aid = create(client, corrected)
    assert client.post(f"/api/student/assignments/{new_aid}/start", headers=headers()).status_code == 200
    assert client.post(f"/api/student/assignments/{old_aid}/start", headers=headers()).status_code == 200
    with db.connection() as conn:
        snapshots = {row["assignment_id"]: json.loads(row["learning_state"])["lesson_snapshot"] for row in conn.execute("SELECT assignment_id,learning_state FROM attempts")}
    assert snapshots[old_aid] == before
    assert snapshots[new_aid]["learning_assets"] == corrected["learning_assets"]
    assert snapshots[old_aid]["learning_assets"] != snapshots[new_aid]["learning_assets"]


def test_targeted_followup_uses_failed_concept_and_handles_revisions(client):
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    answers = [q['correct_index'] for q in plan['diagnostic_quiz']]
    answers[2] = (answers[2] + 1) % 4
    result = quiz(client, aid, plan, answers, 'medium')
    result = turn(client, aid, "Another practice question", action="choice")
    assert result['learning_state']['current_activity']['concept'] == plan['diagnostic_quiz'][2]['concept']
    original = result['learning_state']['diagnostic']
    bad = turn(client, aid, 'It should be quick and nice.')
    assert bad['phase'] == 'adaptive_learning'
    assert bad['learning_state']['current_activity']['revisions'] == 1
    assert bad['learning_state']['demonstrated'] == []
    corrected = turn(client, aid, 'Show the response within two seconds under a load of 300 concurrent requests, and measure each response against that threshold.')
    assert corrected['phase'] == 'adaptive_learning'
    corrected = turn(client, aid, 'The service must return its response in one second at a load of 1000 concurrent requests; record response time and compare with the limit.')
    assert corrected['phase'] == 'reflection'
    assert corrected['learning_state']['diagnostic'] == original


def test_original_keys_are_frozen_in_attempt_snapshot(client):
    from app import db
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    changed = copy.deepcopy(plan)
    changed['diagnostic_quiz'][0]['correct_index'] = (changed['diagnostic_quiz'][0]['correct_index'] + 1) % 4
    with db.connection() as conn:
        conn.execute('UPDATE assignments SET learning_plan=? WHERE id=?', (json.dumps(changed), aid))
    result = quiz(client, aid, plan)
    assert result['quiz_score'] == 5
    assert result['learning_state']['diagnostic'][0]['correct_index'] == plan['diagnostic_quiz'][0]['correct_index']


@pytest.mark.parametrize('provider', ['openai', 'groq'])
def test_generation_preserves_professor_inputs_and_requires_review(client, monkeypatch, provider):
    from app import openai_client
    plan = lesson()
    monkeypatch.setattr(openai_client, 'runtime_ai_mode', lambda: provider)
    monkeypatch.setattr(openai_client, 'GROQ_MODEL', 'openai/gpt-oss-120b')
    def generated(*args, **kwargs):
        assert 'professor-approved' in kwargs['json']['messages'][0]['content']
        assert 'JSON' in kwargs['json']['messages'][0]['content']
        prompt = kwargs['json']['messages'][0]['content']
        assert "refer to concrete details from that question's scenario" in prompt
        assert 'ONE actionable next step or comparison in plain language' in prompt
        assert 'leave the final decision and explanation to them' in prompt
        assert 'pasted unchanged into an unrelated question' in prompt
        assert 'generate reasoning' in prompt
        assert 'not merely entity names' in prompt
        assert 'Do not state the conclusion' in prompt
        assert 'not exact wording or a single keyword' in prompt
        content = kwargs['json']['messages'][1]['content']
        request = json.loads(content)
        assert content == json.dumps(request, ensure_ascii=False, separators=(',', ':'))
        assert request['approved_material'] == plan['approved_material']
        assert ('output_schema' in request) == (provider == 'openai')
        assert len(request['supplemental_reading_excerpts'][0]['excerpt']) == (1000 if provider == 'groq' else 6000)
        output_format = kwargs['json']['response_format']
        assert output_format['type'] == 'json_schema'
        assert output_format['json_schema']['strict'] is True
        schema = output_format['json_schema']['schema']
        assert set(schema['required']) == set(schema['properties'])
        assert 'approved_material' not in schema['properties']
        paraphrased = copy.deepcopy(plan)
        paraphrased['objectives'][0]['description'] = 'A paraphrase produced by the provider.'
        paraphrased['diagnostic_quiz'][3]['reasoning_rubric'] = [criterion['description'] for criterion in plan['diagnostic_quiz'][3]['reasoning_rubric']]
        return httpx.Response(200, request=httpx.Request('POST', 'https://api.openai.com/v1/chat/completions'), json={'choices': [{'message': {'content': json.dumps(paraphrased)}}]})
    monkeypatch.setattr(httpx.Client, 'post', generated)
    inputs = dict(approved_material=plan['approved_material'], objective_descriptions=[o['description'] for o in plan['objectives']], study_resources=plan['study_resources'], intended_difficulty='Introduction', prerequisite_knowledge='None')
    result = openai_client.generate_learning_plan(plan['topic'], 'Undergraduate', study_sources=[{'excerpt': 'x' * 6000}], **inputs)
    assert not result['content_approved']
    assert result['approved_material'] == plan['approved_material']
    assert result['study_resources'] == plan['study_resources']
    assert result['flow_version'] == 2
    assert [o['description'] for o in result['objectives']] == inputs['objective_descriptions']
    monkeypatch.setattr(openai_client, 'runtime_ai_mode', lambda: 'demo')
    with pytest.raises(ValueError, match='unavailable'):
        openai_client.generate_learning_plan(plan['topic'], 'Undergraduate', **inputs)


@pytest.mark.parametrize('cooldown, calls', [('2', 2), ('3600', 1), ('invalid', 1)])
def test_groq_authoring_rate_limit_retry_is_bounded(client, monkeypatch, cooldown, calls):
    from app import openai_client
    monkeypatch.setattr(openai_client, 'runtime_ai_mode', lambda: 'groq')
    sleeps = []
    monkeypatch.setattr(openai_client.time, 'sleep', sleeps.append)
    requests = []
    class Provider:
        def post(self, url, **kwargs):
            requests.append(kwargs)
            return httpx.Response(429, headers={'retry-after': cooldown})
    payload = {'messages': [{'role': 'user', 'content': 'Material unchanged'}]}
    result = openai_client._post_authoring_request(Provider(), 'https://api.groq.com/openai/v1/chat/completions', headers={}, json=payload)
    assert result.status_code == 429
    assert len(requests) == calls
    assert all(request['json'] is payload for request in requests)
    assert len(sleeps) == calls - 1


def test_copying_revealed_quiz_answer_and_misconception_do_not_demonstrate(client):
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    quiz(client, aid, plan)
    turn(client, aid, 'Another practice question', action='choice')
    copied = turn(client, aid, plan['diagnostic_quiz'][0]['explanation'])
    assert copied['phase'] == 'adaptive_learning'
    assert not copied['learning_state']['practice_history'][-1]['demonstrated']
    misconception = turn(client, aid, 'All requirements are functional because every quality constraint is an action.')
    assert misconception['learning_state']['practice_history'][-1]['category'] == 'identifiable_misconception'
    assert 'counterexample' in misconception['messages'][-1]['content'].lower()


@pytest.mark.parametrize('choice,marker,supported,level', [
    ('A simpler explanation', "Let's break down", True, 'foundational'),
    ('A worked example', 'Worked example', True, 'foundational'),
    ('Another practice question', 'Fresh practice', False, 'standard'),
    ('An application challenge', 'Application challenge', False, 'accelerated'),
    ('A recap', 'Recap', True, 'standard'),
])
def test_each_post_quiz_choice_has_distinct_behavior_and_disappears(client, choice, marker, supported, level):
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    answers = [q['correct_index'] for q in plan['diagnostic_quiz']]
    answers[0] = (answers[0] + 1) % 4
    result = quiz(client, aid, plan, answers, 'medium')
    assert result['phase'] == 'learning_choice'
    assert 'current_activity' not in result['learning_state']
    assert len(result['learning_state']['choices']) == 5
    assert set(result['learning_state']['choice_descriptions']) == set(result['learning_state']['choices'])
    assert 'Functional behavior' in result['messages'][-1]['content']
    selected = turn(client, aid, choice, action='choice')
    assert selected['phase'] == 'adaptive_learning'
    assert marker in selected['messages'][-1]['content']
    assert selected['learning_state']['choices'] == []
    assert selected['learning_state']['current_activity']['assisted'] is supported
    assert selected['learning_state']['current_activity']['level'] == level
    turn(client, aid, action='pause')
    resumed = turn(client, aid, action='resume')
    assert resumed['learning_state']['choices'] == []
    unknown_reply = turn(client, aid, "I don't know")
    assert unknown_reply['learning_state']['choices'] == []
    assert 'smaller step' in unknown_reply['messages'][-1]['content']
    revision = turn(client, aid, 'I think it is just good software.')
    assert revision['phase'] == 'adaptive_learning'
    assert revision['learning_state']['choices'] == []
    assert revision['learning_state']['current_activity']['revisions'] == 1


def test_existing_post_quiz_session_receives_chooser_without_repeating_quiz(client):
    from app import db
    plan = lesson()
    aid = create(client, plan)
    start_quiz(client, aid)
    completed = quiz(client, aid, plan)
    with db.connection() as conn:
        row = conn.execute('SELECT * FROM attempts WHERE assignment_id=? AND student_id=?', (aid, 'student-alex')).fetchone()
        saved = json.loads(row['learning_state'])
        saved.pop('choice_flow_version', None)
        saved['stage'] = 'adaptive_learning'
        conn.execute('UPDATE attempts SET phase=?,learning_state=? WHERE id=?', ('adaptive_learning', json.dumps(saved), row['id']))
    reopened = client.post(f'/api/student/assignments/{aid}/start', headers=headers()).json()
    assert reopened['phase'] == 'learning_choice'
    assert reopened['learning_state']['diagnostic'] == completed['learning_state']['diagnostic']
    assert reopened['quiz_score'] == completed['quiz_score']
    again = client.post(f'/api/student/assignments/{aid}/start', headers=headers()).json()
    assert len(again['messages']) == len(reopened['messages'])


def test_production_service_rejects_direct_demo_identity(client, monkeypatch):
    monkeypatch.setenv('ADAPTIVE_PLATFORM_ONLY','true')
    monkeypatch.setenv('PLATFORM_SERVICE_TOKEN','test-service-token')
    assert client.get('/api/student/assignments',headers=headers()).status_code == 401
    assert client.get('/api/demo/users').status_code == 401
    response=client.get('/api/student/assignments',headers={**headers(),'X-Platform-Service':'test-service-token'})
    assert response.status_code == 200
    assert client.get('/api/health').status_code == 200
