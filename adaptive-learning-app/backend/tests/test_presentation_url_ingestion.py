"""Real presentation/article parsing; bounded mocked HTTP and existing RAG storage."""
import io
import json
import socket
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from pptx import Presentation
from pptx.util import Inches

from test_api import client, create, headers
from test_pdf_ingestion import make_pdf


def presentation(*, slides=2, notes=False, table=False, oversized=False):
    deck = Presentation()
    for number in range(1, slides + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Process States {number}"
        frame = slide.placeholders[1].text_frame
        frame.text = "Ready waits for CPU"
        paragraph = frame.add_paragraph()
        paragraph.text = "Running executes instructions"
        paragraph.level = 1
        if oversized:
            paragraph = frame.add_paragraph()
            paragraph.text = "A blocked process waits for input and releases the CPU. " * 150
        if notes:
            slide.notes_slide.notes_text_frame.text = "Explain why blocked processes release the CPU."
        if table:
            grid = slide.shapes.add_table(3, 3, Inches(1), Inches(4), Inches(7), Inches(1)).table
            for row, values in zip(grid.rows, [("State", "Description", "CPU"),
                                               ("Running", "Executing instructions", "Yes"),
                                               ("Blocked", "Waiting for input", "No")]):
                for cell, value in zip(row.cells, values):
                    cell.text = value
    result = io.BytesIO()
    deck.save(result)
    return result.getvalue()


def parse_deck(tmp_path, **options):
    from app.resource_ingestion import ingest_resource
    path = tmp_path / "lecture.pptx"
    path.write_bytes(presentation(**options))
    return ingest_resource(path)


def test_presentation_keeps_slides_titles_and_bullet_hierarchy(tmp_path):
    chunks = parse_deck(tmp_path)
    assert len(chunks) == 2
    for index, chunk in enumerate(chunks, 1):
        assert chunk.position == index - 1
        assert chunk.metadata == {"source_type": "pptx", "slide": index, "title": f"Process States {index}"}
        assert "- Ready waits for CPU" in chunk.content
        assert "  - Running executes instructions" in chunk.content
        assert f"Process States {3 - index}" not in chunk.content


def test_presentation_tables_and_existing_notes(tmp_path):
    chunks = parse_deck(tmp_path, slides=1, table=True, notes=True)
    assert len(chunks) == 1
    assert "State | Description | CPU" in chunks[0].content
    assert "Running | Executing instructions | Yes" in chunks[0].content
    assert "Blocked | Waiting for input | No" in chunks[0].content
    assert "Speaker Notes:\nExplain why blocked processes release the CPU." in chunks[0].content


def test_parser_does_not_access_missing_notes(tmp_path, monkeypatch):
    from pptx.slide import Slide
    monkeypatch.setattr(Slide, "notes_slide", property(lambda self: pytest.fail("Created a missing notes slide")))
    assert parse_deck(tmp_path)


def test_oversized_slide_splits_with_context_and_budget(tmp_path, monkeypatch):
    import tiktoken
    from app import resource_ingestion as ingestion
    monkeypatch.setattr(ingestion, "RESOURCE_CHUNK_MAX_TOKENS", 90)
    chunks = parse_deck(tmp_path, oversized=True)
    assert len(chunks) > 2
    for chunk in chunks:
        assert f"Slide {chunk.metadata['slide']}\nTitle: {chunk.metadata['title']}" in chunk.content
        assert len(tiktoken.encoding_for_model(ingestion.OPENAI_EMBEDDING_MODEL).encode(chunk.content)) <= 90
        assert f"Process States {3 - chunk.metadata['slide']}" not in chunk.content


def test_legacy_ppt_missing_converter_is_clear_and_atomic(client, monkeypatch):
    from app import resource_ingestion as ingestion, db, config
    monkeypatch.setattr(ingestion.shutil, "which", lambda name: None)
    monkeypatch.setattr(ingestion.Path, "is_file", lambda path: False)
    aid = create(client, publish=False)
    result = client.post(f"/api/instructor/assignments/{aid}/documents", headers=headers("instructor-demo"),
                         files={"file": ("old.ppt", b"legacy binary")})
    assert result.status_code == 503
    assert "LibreOffice" in result.json()["detail"] and "PPTX" in result.json()["detail"]
    with db.connection() as conn:
        assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 0
    assert not list(config.UPLOAD_DIR.iterdir())


def test_legacy_conversion_routes_to_pptx_and_cleans_temp(tmp_path, monkeypatch):
    from app import resource_ingestion as ingestion
    monkeypatch.setattr(ingestion.shutil, "which", lambda name: "soffice")
    seen = []
    def convert(command, **kwargs):
        directory = Path(command[command.index("--outdir") + 1])
        seen.append(directory)
        assert "--headless" in command and kwargs["timeout"] == 60
        (directory / "old.pptx").write_bytes(presentation(slides=1))
    monkeypatch.setattr(ingestion.subprocess, "run", convert)
    path = tmp_path / "old.ppt"
    path.write_bytes(b"legacy binary")
    assert ingestion.ingest_resource(path)[0].metadata["slide"] == 1
    assert not seen[0].exists()


ARTICLE = b'''<html><head><title>Process Management</title></head><body>
<nav>Navigation home contact unrelated links</nav><article><h1>Process Management</h1>
<p>A process is a running instance of a program. The operating system schedules each process fairly on the CPU.</p>
<h2>Ready processes</h2><p>A ready process waits for CPU time. Scheduling lets many applications share a processor and keeps them responsive.</p>
<ul><li>Ready: waiting for the CPU.</li><li>Running: executing instructions.</li></ul>
<h2>Blocked processes</h2><p>A blocked process waits for input and releases the CPU. When input arrives it can become ready again.</p>
<table><tr><th>State</th><th>CPU</th></tr><tr><td>Running</td><td>Yes</td></tr><tr><td>Blocked</td><td>No</td></tr></table>
</article><footer>Unrelated footer advertisement</footer></body></html>'''


@pytest.fixture
def network(monkeypatch):
    from app import resource_ingestion as ingestion
    real_client = httpx.Client
    calls = []
    routes = {}
    def resolve(host, port, **kwargs):
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("127.0.0.1" if host == "private.test" else "93.184.216.34", port))]
    monkeypatch.setattr(ingestion, "socket", SimpleNamespace(getaddrinfo=resolve, SOCK_STREAM=socket.SOCK_STREAM))
    def request(req):
        calls.append(str(req.url))
        result = routes[str(req.url)]
        if isinstance(result, Exception):
            raise result
        return result
    # Isolate resource HTTP mocks from Docling's model downloads and TestClient.
    class MockHTTPClient(real_client):
        def __init__(self, **kwargs):
            super().__init__(transport=httpx.MockTransport(request), **kwargs)
    monkeypatch.setattr(ingestion, "httpx", SimpleNamespace(Client=MockHTTPClient, HTTPError=httpx.HTTPError,
                                                          TimeoutException=httpx.TimeoutException))
    return routes, calls


def test_article_retains_structure_and_discards_chrome(tmp_path):
    from app.resource_ingestion import ingest_html
    path = tmp_path / "article.html"
    path.write_bytes(ARTICLE)
    chunks = ingest_html(path, "https://school.test/start", "https://school.test/article")
    combined = "\n".join(chunk.content for chunk in chunks)
    assert "Navigation home" not in combined and "footer advertisement" not in combined
    assert "Ready processes" in combined and "Blocked processes" in combined
    assert "State | CPU" in combined and "Running | Yes" in combined
    assert any("- Ready:" in chunk.content for chunk in chunks)
    assert any("Ready processes" in chunk.metadata.get("section", "") for chunk in chunks)
    assert all(chunk.metadata["source_url"] == "https://school.test/start" for chunk in chunks)


def test_large_html_section_repeats_context_and_respects_tokens(tmp_path, monkeypatch):
    import tiktoken
    from app import resource_ingestion as ingestion
    monkeypatch.setattr(ingestion, "RESOURCE_CHUNK_MAX_TOKENS", 80)
    path = tmp_path / "article.html"
    path.write_bytes(ARTICLE.replace(b"A ready process waits for CPU time.",
                                    b"A ready process waits for CPU time. " * 100))
    chunks = ingestion.ingest_html(path, "https://school.test/article", "https://school.test/article")
    ready_chunks = [chunk for chunk in chunks if "Ready processes" in chunk.metadata.get("section", "")]
    assert len(ready_chunks) > 1
    tokenizer = tiktoken.encoding_for_model(ingestion.OPENAI_EMBEDDING_MODEL)
    assert all(len(tokenizer.encode(chunk.content)) <= 80 for chunk in chunks)
    assert all("Process Management" in chunk.content and "Ready processes" in chunk.content for chunk in ready_chunks)


def test_invalid_pptx_is_rejected_and_removed(client):
    from app import db, config
    aid = create(client, publish=False)
    result = client.post(f"/api/instructor/assignments/{aid}/documents", headers=headers("instructor-demo"), files={"file": ("bad.pptx", b"not a presentation")})
    assert result.status_code == 422 and "presentation" in result.json()["detail"]
    with db.connection() as conn:
        assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 0
    assert not list(config.UPLOAD_DIR.iterdir())


def test_safe_redirect_retains_original_and_final_url(client, network):
    from app import db
    routes, calls = network
    routes["https://school.test/start"] = httpx.Response(302, headers={"location": "/article"})
    routes["https://school.test/article"] = httpx.Response(200, content=ARTICLE, headers={"content-type": "text/html"})
    aid = create(client, publish=False)
    result = client.post(f"/api/instructor/assignments/{aid}/documents/url", headers=headers("instructor-demo"), json={"url": "https://school.test/start"})
    assert result.status_code == 201, result.text
    assert calls == ["https://school.test/start", "https://school.test/article"]
    with db.connection() as conn:
        metadata = json.loads(conn.execute("SELECT metadata FROM chunks LIMIT 1").fetchone()[0])
    assert metadata["source_url"] == "https://school.test/start"
    assert metadata["final_url"] == "https://school.test/article"


@pytest.mark.parametrize("url", ["file:///private", "http://private.test/", "https://youtube.com/watch?v=123", "https://instagram.com/post", "https://user:secret@school.test/", "https://school.test:9999/"])
def test_invalid_urls_never_fetch_or_save(client, network, url):
    from app import db, config
    _, calls = network
    aid = create(client, publish=False)
    result = client.post(f"/api/instructor/assignments/{aid}/documents/url", headers=headers("instructor-demo"), json={"url": url})
    assert result.status_code == 422, result.text
    assert not calls
    with db.connection() as conn:
        assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 0
    assert not list(config.UPLOAD_DIR.iterdir())


@pytest.mark.parametrize("scenario", ["private_redirect", "too_many", "binary", "oversized", "timeout", "private_page", "empty_article"])
def test_url_failures_leave_no_records_or_files(client, network, scenario, monkeypatch):
    from app import resource_ingestion as ingestion, db, config
    routes, calls = network
    url = "https://school.test/resource"
    cases = {
        "private_redirect": httpx.Response(302, headers={"location": "http://private.test/"}),
        "too_many": httpx.Response(302, headers={"location": url}),
        "binary": httpx.Response(200, content=b"not a PDF", headers={"content-type": "application/octet-stream"}),
        "oversized": httpx.Response(200, content=b"x" * 100, headers={"content-type": "text/html"}),
        "timeout": httpx.ReadTimeout("timed out"),
        "private_page": httpx.Response(403),
        "empty_article": httpx.Response(200, content=b"<html></html>", headers={"content-type": "text/html"}),
    }
    routes[url] = cases[scenario]
    if scenario == "oversized":
        monkeypatch.setattr(ingestion, "RESOURCE_URL_MAX_BYTES", 30)
    aid = create(client, publish=False)
    result = client.post(f"/api/instructor/assignments/{aid}/documents/url", headers=headers("instructor-demo"), json={"url": url})
    assert result.status_code == 422, result.text
    assert "http://private.test/" not in calls
    with db.connection() as conn:
        assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM chunks").fetchone()[0] == 0
    assert not list(config.UPLOAD_DIR.iterdir())


@pytest.mark.parametrize("name,mime", [("paper.pdf", "application/pdf"), ("download", "application/octet-stream")])
def test_pdf_urls_use_existing_pdf_function(client, network, monkeypatch, name, mime):
    from app import resource_ingestion as ingestion, db
    routes, _ = network
    url = f"https://school.test/{name}"
    routes[url] = httpx.Response(200, content=make_pdf(), headers={"content-type": mime})
    actual_pdf = ingestion.ingest_pdf
    seen = []
    def parse(path):
        seen.append(path)
        return actual_pdf(path)
    monkeypatch.setattr(ingestion, "ingest_pdf", parse)
    aid = create(client, publish=False)
    result = client.post(f"/api/instructor/assignments/{aid}/documents/url", headers=headers("instructor-demo"), json={"url": url})
    assert result.status_code == 201, result.text
    assert len(seen) == 1 and seen[0].suffix == ".pdf"
    with db.connection() as conn:
        rows = conn.execute("SELECT metadata FROM chunks").fetchall()
    assert all(json.loads(row[0])["pages"] == [1] and json.loads(row[0])["source_url"] == url for row in rows)


@pytest.mark.parametrize("kind", ["pptx", "html"])
def test_complete_pipeline_pairing_isolation_and_top_four(client, network, monkeypatch, kind):
    from app import rag, db, resource_ingestion as ingestion
    monkeypatch.setattr(ingestion, "RESOURCE_CHUNK_MAX_TOKENS", 60)
    embedded = []
    def embed(texts):
        embedded.extend(texts)
        return [[index + 1.0, 1.0] for index in range(len(texts))]
    monkeypatch.setattr(rag, "embed_texts", embed)
    aid = create(client, publish=False)
    other = create(client, publish=False)
    if kind == "pptx":
        result = client.post(f"/api/instructor/assignments/{aid}/documents", headers=headers("instructor-demo"), files={"file": ("lecture.pptx", presentation(slides=6))})
    else:
        routes, _ = network
        extended = ARTICLE.replace(b"</article>", b"".join(
            f"<h2>Scheduling strategy {index}</h2><p>Scheduling strategy {index} considers ready processes and allocates CPU time fairly to responsive applications.</p>".encode()
            for index in range(6)) + b"</article>")
        routes["https://school.test/article"] = httpx.Response(200, content=extended, headers={"content-type": "text/html"})
        result = client.post(f"/api/instructor/assignments/{aid}/documents/url", headers=headers("instructor-demo"), json={"url": "https://school.test/article"})
    assert result.status_code == 201, result.text
    with db.connection() as conn:
        rows = conn.execute("SELECT content,embedding,metadata,position FROM chunks ORDER BY position").fetchall()
    assert len(rows) > 4
    assert [row[0] for row in rows] == embedded
    assert [json.loads(row[1]) for row in rows] == [[index + 1.0, 1.0] for index in range(len(rows))]
    assert all(row[0] and json.loads(row[2])["source_type"] == kind for row in rows)
    assert [row[3] for row in rows] == list(range(len(rows)))
    monkeypatch.setattr(rag, "embed_texts", lambda texts: [[1.0, 1.0] for _ in texts])
    retrieved = rag.retrieve(aid, "process")
    assert len(retrieved) == 4
    assert [item["content"] for item in retrieved] == [row[0] for row in rows[:4]]
    assert all(item["document_id"] == result.json()["id"] for item in retrieved)
    assert rag.retrieve(other, "process") == []


def test_url_auth_and_draft_restrictions(client, network):
    _, calls = network
    aid = create(client, publish=False)
    endpoint = f"/api/instructor/assignments/{aid}/documents/url"
    assert client.post(endpoint, headers=headers(), json={"url": "https://school.test/"}).status_code == 403
    published = create(client)
    assert client.post(f"/api/instructor/assignments/{published}/documents/url", headers=headers("instructor-demo"), json={"url": "https://school.test/"}).status_code == 409
    assert not calls


@pytest.mark.parametrize("failure,status", [("embedding", 502), ("storage", 500)])
def test_url_indexing_failure_is_atomic(client, network, monkeypatch, failure, status):
    from app import rag, db, config
    routes, _ = network
    url = "https://school.test/article"
    routes[url] = httpx.Response(200, content=ARTICLE, headers={"content-type": "text/html"})
    aid = create(client, publish=False)
    if failure == "embedding":
        monkeypatch.setattr(rag, "embed_texts", lambda texts: (_ for _ in ()).throw(RuntimeError("provider error")))
    else:
        # Strict pairing fails after INSERT documents; the transaction must roll back.
        monkeypatch.setattr(rag, "embed_texts", lambda texts: [])
    result = client.post(f"/api/instructor/assignments/{aid}/documents/url", headers=headers("instructor-demo"), json={"url": url})
    assert result.status_code == status, result.text
    with db.connection() as conn:
        assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 0
    assert not list(config.UPLOAD_DIR.iterdir())


def test_stream_size_is_bounded_without_content_length(network, monkeypatch):
    from app import resource_ingestion as ingestion
    routes, _ = network
    monkeypatch.setattr(ingestion, "RESOURCE_URL_MAX_BYTES", 30)
    routes["https://school.test/article"] = httpx.Response(200, stream=httpx.ByteStream(b"x" * 100), headers={"content-type": "text/html"})
    with pytest.raises(ingestion.ResourceIngestionError, match="smaller"):
        ingestion.fetch_resource_url("https://school.test/article")


def test_html_url_with_pdf_filename_is_not_routed_to_pdf(client, network, monkeypatch):
    from app import resource_ingestion as ingestion
    routes, _ = network
    url = "https://school.test/article.pdf"
    routes[url] = httpx.Response(200, content=ARTICLE, headers={"content-type": "text/html"})
    monkeypatch.setattr(ingestion, "ingest_pdf", lambda path: pytest.fail("HTML filename must not select PDF parser"))
    aid = create(client, publish=False)
    result = client.post(f"/api/instructor/assignments/{aid}/documents/url", headers=headers("instructor-demo"), json={"url": url})
    assert result.status_code == 201, result.text


def test_pptx_preview_uses_slide_parser(client):
    result = client.post("/api/instructor/documents/preview", headers=headers("instructor-demo"), files={"file": ("lecture.pptx", presentation(notes=True))})
    assert result.status_code == 200, result.text
    assert "Slide 1" in result.json()["text"] and "Slide 2" in result.json()["text"]
    assert "Speaker Notes:" in result.json()["text"]
