"""Format-aware parsing before the existing RAG embedding/storage pipeline.

PDF uses Docling; presentations preserve slides; URLs use article extraction
or the same PDF path. TXT/Markdown keep their existing rag.py path.
"""
from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from threading import Lock
import ipaddress
import shutil
import socket
import subprocess
import time
from tempfile import TemporaryDirectory
from urllib.parse import urljoin, urlsplit
from xml.etree import ElementTree

import httpx

from .config import (OPENAI_EMBEDDING_MODEL, PDF_CHUNK_MAX_TOKENS, PDF_MODEL_ARTIFACTS_PATH,
                     RESOURCE_CHUNK_MAX_TOKENS, RESOURCE_URL_MAX_BYTES, RESOURCE_URL_TIMEOUT_SECONDS)


class ResourceIngestionError(ValueError):
    """An expected resource failure whose explanation can be shown to instructors."""


class PowerPointConversionUnavailable(ResourceIngestionError):
    pass


@dataclass
class ResourceChunk:
    """Parser output: exact embedding text, zero-based position, JSON-safe metadata.

    Storage supplies document_id; parsing does not embed or persist this object.
    """
    content: str
    position: int
    metadata: dict


# Reuse expensive PDF models; serialize conversion on this shared instance.
_conversion_lock = Lock()


@lru_cache(maxsize=1)
def _pdf_converter():
    from docling.backend.pypdfium2_backend import PyPdfiumDocumentBackend
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import OcrMode, PdfPipelineOptions, RapidOcrOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    options = PdfPipelineOptions(
        do_table_structure=True,
        do_ocr=True,
        ocr_options=RapidOcrOptions(mode=OcrMode.PDF_AWARE_LAYOUT_REGIONS),
        artifacts_path=PDF_MODEL_ARTIFACTS_PATH or None,
    )
    return DocumentConverter(
        allowed_formats=[InputFormat.PDF],
        # Use Docling's supported PDFium backend: the default threaded native
        # renderer emitted Windows access violations on the scanned fixture.
        format_options={InputFormat.PDF: PdfFormatOption(
            pipeline_options=options, backend=PyPdfiumDocumentBackend,
        )},
    )


def ingest_resource(path: Path) -> list[ResourceChunk]:
    """Route supported resources without embedding, persistence, or lesson logic."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return ingest_pdf(path)
    if suffix == ".pptx":
        return ingest_pptx(path)
    if suffix == ".ppt":
        return ingest_ppt(path)
    raise ValueError("Structured ingestion supports PDF, PPTX, and PPT.")


def ingest_pdf(path: Path) -> list[ResourceChunk]:
    import tiktoken
    from docling.chunking import HybridChunker
    from docling.datamodel.base_models import ConversionStatus
    from docling_core.transforms.chunker.tokenizer.openai import OpenAITokenizer
    from docling_core.types.doc import DocItemLabel

    if PDF_CHUNK_MAX_TOKENS < 1:
        raise ValueError("PDF_CHUNK_MAX_TOKENS must be positive.")
    with _conversion_lock:
        result = _pdf_converter().convert(path)
    if result.status != ConversionStatus.SUCCESS:
        raise ValueError("The PDF could not be fully processed.")
    document = result.document

    # A filename is not a document title: include only an extracted title.
    title = next((item.text.strip() for item, _ in document.iterate_items()
                  if item.label == DocItemLabel.TITLE and getattr(item, "text", "").strip()), None)
    tokenizer = OpenAITokenizer(
        tokenizer=tiktoken.encoding_for_model(OPENAI_EMBEDDING_MODEL),
        max_tokens=PDF_CHUNK_MAX_TOKENS,
    )
    chunker = HybridChunker(tokenizer=tokenizer, repeat_table_header=True)
    chunks = []
    for chunk in chunker.chunk(dl_doc=document):
        content = chunker.contextualize(chunk=chunk).strip()
        if not chunk.text.strip() or not content:
            continue
        metadata = {"source_type": "pdf"}
        # Docling provenance uses one-based physical pages, including all pages
        # contributing to a chunk rather than just its first page.
        pages = sorted({prov.page_no for item in chunk.meta.doc_items
                        for prov in item.prov if prov.page_no >= 1})
        if pages:
            metadata["pages"] = pages
        if title:
            metadata["title"] = title
        if chunk.meta.headings:
            metadata["section"] = " / ".join(chunk.meta.headings)
        chunks.append(ResourceChunk(content, len(chunks), metadata))
    if not chunks:
        raise ValueError("No readable PDF content was found.")
    return chunks


# Presentation and HTML chunking deliberately do not use PDF HybridChunker.
def _chunk_blocks(context: str, blocks: list[tuple[str, str]], budget: int) -> list[str]:
    """Pack semantic blocks, repeating context and labels when a block is split."""
    import tiktoken
    tokenizer = tiktoken.encoding_for_model(OPENAI_EMBEDDING_MODEL)
    count = lambda value: len(tokenizer.encode(value, disallowed_special=()))
    if budget < 1 or count(context + "\n\n") >= budget:
        raise ResourceIngestionError("The chunk token budget is too small for the resource title/context.")
    result, current = [], context
    for label, text in blocks:
        remaining = text.strip()
        if not remaining:
            continue
        if count(context + "\n\n" + label) >= budget:
            raise ResourceIngestionError("The chunk token budget is too small for a table header or section label.")
        while remaining:
            candidate = current + "\n\n" + label + remaining
            if count(candidate) <= budget:
                current = candidate
                break
            if current != context:
                result.append(current)
                current = context
                continue
            # Find a token-bounded prefix, preferably ending at a word boundary.
            low, high = 0, len(remaining)
            while low < high:
                middle = (low + high + 1) // 2
                if count(context + "\n\n" + label + remaining[:middle]) <= budget:
                    low = middle
                else:
                    high = middle - 1
            if not low:
                raise ResourceIngestionError("The chunk token budget cannot fit the resource content.")
            boundary = remaining.rfind(" ", 0, low + 1)
            end = boundary if boundary > 0 else low
            result.append(context + "\n\n" + label + remaining[:end].rstrip())
            remaining = remaining[end:].lstrip()
    if current != context:
        result.append(current)
    return result


def _table_blocks(rows: list[list[str]]) -> list[tuple[str, str]]:
    if not rows:
        return []
    header = "Table:\n" + " | ".join(rows[0]) + "\n"
    return [(header, " | ".join(row)) for row in rows[1:]] or [("Table:\n", " | ".join(rows[0]))]


def ingest_pptx(path: Path) -> list[ResourceChunk]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    presentation = Presentation(path)
    chunks = []

    def shapes_in_order(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shapes_in_order(shape.shapes)
            else:
                yield shape

    for number, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        context = f"Slide {number}" + (f"\nTitle: {title}" if title else "")
        blocks = []
        for shape in shapes_in_order(slide.shapes):
            if shape is title_shape or (title_shape is not None and shape.shape_id == title_shape.shape_id):
                continue
            if shape.has_table:
                blocks.extend(_table_blocks([[cell.text.strip().replace("\n", " / ") for cell in row.cells]
                                             for row in shape.table.rows]))
            elif shape.has_text_frame:
                blocks.extend(("  " * paragraph.level + "- ", paragraph.text)
                              for paragraph in shape.text_frame.paragraphs if paragraph.text.strip())
        # Accessing notes_slide without this check would create missing notes.
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                blocks.append(("Speaker Notes:\n", notes.text))
        if not blocks and title:
            blocks = [("", title)]
        metadata = {"source_type": "pptx", "slide": number}
        if title:
            metadata["title"] = title
        for content in _chunk_blocks(context, blocks, RESOURCE_CHUNK_MAX_TOKENS):
            chunks.append(ResourceChunk(content, len(chunks), dict(metadata)))
    if not chunks:
        raise ResourceIngestionError("No readable slide text was found in this presentation.")
    return chunks


def ingest_ppt(path: Path) -> list[ResourceChunk]:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        for candidate in (Path("C:/Program Files/LibreOffice/program/soffice.exe"),
                          Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe")):
            if candidate.is_file():
                executable = str(candidate)
                break
    if not executable:
        raise PowerPointConversionUnavailable("Legacy PPT requires LibreOffice, which is not installed. Export as PPTX and upload again.")
    with TemporaryDirectory() as directory:
        root = Path(directory)
        # A separate profile prevents collisions with an open desktop instance.
        try:
            subprocess.run([executable, f"-env:UserInstallation={(root / 'profile').as_uri()}",
                            "--headless", "--convert-to", "pptx", "--outdir", str(root), str(path.resolve())],
                           check=True, timeout=60, capture_output=True)
        except (subprocess.SubprocessError, OSError) as exc:
            raise ResourceIngestionError("LibreOffice could not convert this PPT. Export as PPTX and upload again.") from exc
        converted = root / (path.stem + ".pptx")
        if not converted.is_file():
            raise ResourceIngestionError("LibreOffice did not produce a PPTX. Export as PPTX and upload again.")
        return ingest_pptx(converted)


_UNSUPPORTED_URL_HOSTS = {"youtube.com", "youtu.be", "vimeo.com", "dailymotion.com", "twitch.tv",
                          "tiktok.com", "facebook.com", "instagram.com", "twitter.com", "x.com",
                          "reddit.com", "linkedin.com"}


def _validate_resource_url(url: str) -> None:
    try:
        parsed = urlsplit(url)
        if len(url) > 2000 or parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
            raise ResourceIngestionError("Provide a public HTTP or HTTPS article/PDF URL without login credentials.")
        host = parsed.hostname.lower().rstrip(".")
        if any(host == blocked or host.endswith("." + blocked) for blocked in _UNSUPPORTED_URL_HOSTS):
            raise ResourceIngestionError("Video and social-media URLs are not supported. Use an article or PDF.")
        if parsed.port not in {None, 80, 443}:
            raise ResourceIngestionError("Only standard public HTTP/HTTPS ports are supported.")
        addresses = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)
        if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
            raise ResourceIngestionError("Private and internal network URLs are not supported.")
    except ResourceIngestionError:
        raise
    except (ValueError, OSError) as exc:
        raise ResourceIngestionError("The resource URL is invalid or its public host could not be resolved.") from exc


def fetch_resource_url(url: str) -> tuple[bytes, str, str]:
    """Fetch one explicit resource, validating every redirect and bounding reads."""
    current = url.strip()
    deadline = time.monotonic() + RESOURCE_URL_TIMEOUT_SECONDS
    try:
        with httpx.Client(timeout=RESOURCE_URL_TIMEOUT_SECONDS, follow_redirects=False, trust_env=False) as client:
            for _ in range(6):
                _validate_resource_url(current)
                if time.monotonic() >= deadline:
                    raise ResourceIngestionError("The resource download timed out.")
                remaining_time = deadline - time.monotonic()
                if remaining_time <= 0:
                    raise ResourceIngestionError("The resource download timed out.")
                with client.stream("GET", current, timeout=remaining_time,
                                   headers={"Accept": "text/html,application/pdf", "User-Agent": "AdaptiveLearningResource/1.0"}) as response:
                    if time.monotonic() >= deadline:
                        raise ResourceIngestionError("The resource download timed out.")
                    if response.status_code in {301, 302, 303, 307, 308}:
                        location = response.headers.get("location")
                        if not location:
                            raise ResourceIngestionError("The resource redirect has no destination.")
                        current = urljoin(current, location)
                        continue
                    if response.status_code != 200:
                        raise ResourceIngestionError("The resource must be a publicly accessible article or PDF (HTTP 200).")
                    if int(response.headers.get("content-length", "0")) > RESOURCE_URL_MAX_BYTES:
                        raise ResourceIngestionError("URL resources must be 10 MB or smaller.")
                    data = bytearray()
                    # Inspect each transport read. Rechunking into 64 KB blocks
                    # can hide a slow trickle from the elapsed-time checks.
                    for part in response.iter_bytes():
                        if time.monotonic() >= deadline:
                            raise ResourceIngestionError("The resource download timed out.")
                        data.extend(part)
                        if len(data) > RESOURCE_URL_MAX_BYTES:
                            raise ResourceIngestionError("URL resources must be 10 MB or smaller.")
                    mime = response.headers.get("content-type", "").split(";")[0].strip().lower()
                    if bytes(data[:1024]).lstrip().startswith(b"%PDF-"):
                        return bytes(data), "application/pdf", current
                    if mime in {"text/html", "application/xhtml+xml"} and data:
                        return bytes(data), "text/html", current
                    raise ResourceIngestionError("Only HTML articles and actual PDF documents are supported.")
            raise ResourceIngestionError("The resource redirected too many times.")
    except ResourceIngestionError:
        raise
    except httpx.TimeoutException as exc:
        raise ResourceIngestionError("The resource download timed out.") from exc
    except (httpx.HTTPError, ValueError) as exc:
        raise ResourceIngestionError("The resource could not be downloaded. Check the public URL and try again.") from exc


def ingest_html(path: Path, source_url: str, final_url: str) -> list[ResourceChunk]:
    import trafilatura
    extracted = trafilatura.extract(path.read_bytes(), url=final_url, output_format="xml",
                                    with_metadata=True, include_comments=False,
                                    include_tables=True, include_formatting=True, favor_precision=True)
    if not extracted:
        raise ResourceIngestionError("No readable article content was found at this URL.")
    document = ElementTree.fromstring(extracted)
    title = document.get("title", "").strip()
    chunks, blocks, headings = [], [], {}

    def flush():
        section = " / ".join(headings.values())
        context = "\n".join(value for value in (title, section) if value) or "Article"
        metadata = {"source_type": "html", "source_url": source_url, "final_url": final_url}
        if title:
            metadata["title"] = title
        if section:
            metadata["section"] = section
        for content in _chunk_blocks(context, blocks, RESOURCE_CHUNK_MAX_TOKENS):
            chunks.append(ResourceChunk(content, len(chunks), dict(metadata)))
        blocks.clear()

    def text(node):
        return " ".join(" ".join(node.itertext()).split())

    def walk(node, depth=0):
        for child in node:
            if child.tag == "head":
                flush()
                level = int(child.get("rend", "h2").lstrip("h"))
                for existing in list(headings):
                    if existing >= level:
                        del headings[existing]
                headings[level] = text(child)
            elif child.tag == "table":
                blocks.extend(_table_blocks([[text(cell) for cell in row] for row in child.findall("row")]))
            elif child.tag == "list":
                walk(child, depth + 1)
            elif child.tag == "item":
                # Nested list text is handled separately to retain its hierarchy.
                value = " ".join([child.text or ""] + [text(part) + " " + (part.tail or "")
                                  for part in child if part.tag != "list"]).strip()
                blocks.append(("  " * max(0, depth - 1) + "- ", value))
                for nested in child.findall("list"):
                    walk(nested, depth + 1)
            elif child.tag in {"p", "quote", "code"}:
                blocks.append(("", text(child)))
            else:
                walk(child, depth)
    main = document.find("main")
    if main is not None:
        walk(main)
    flush()
    if not chunks:
        raise ResourceIngestionError("No readable article content was found at this URL.")
    return chunks
